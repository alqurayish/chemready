"""End-to-end tests for the backend API, using the rules baseline (no AI, no network)."""

import io
import time
from datetime import date
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from chemready.app import auth
from chemready.app.main import create_app
from chemready.app.service import Worker
from chemready.app.store import Store
from chemready.config import Settings
from chemready.demo.synthetic import make_sds_pdf
from chemready.demo.synthetic_set import SPECS
from chemready.extraction.baseline import RulesClient
from chemready.extraction.llm import FakeClient

PASSWORD = "correct horse battery"


def _pdf(sds_id: str) -> bytes:
    return make_sds_pdf(SPECS[sds_id][0])


@pytest.fixture
def settings(tmp_path: Path) -> Settings:
    return Settings(
        database_path=tmp_path / "db.sqlite", upload_dir=tmp_path / "uploads", llm_provider="rules"
    )


@pytest.fixture
def client(settings: Settings) -> TestClient:
    app = create_app(settings, client_factory=RulesClient, inline_processing=True)
    app.state.chemready.today = lambda: date(2026, 10, 9)
    return TestClient(app)


def _sign_up(client: TestClient, email: str = "rahima@demo.test", facility: str = "Demo Washing Ltd") -> None:
    response = client.post(
        "/api/auth/signup",
        json={"name": "Rahima Akter", "email": email, "password": PASSWORD, "facility_name": facility},
    )
    assert response.status_code == 201, response.text


def _upload(client: TestClient, *files: tuple[str, bytes], public: bool = False) -> list[dict[str, object]]:
    response = client.post(
        "/api/documents",
        files=[("files", (name, data, "application/pdf")) for name, data in files],
        data={"public_sds": str(public).lower()},
    )
    assert response.status_code == 200, response.text
    results: list[dict[str, object]] = response.json()["results"]
    return results


def test_protected_routes_need_sign_in(client: TestClient) -> None:
    for path in ("/api/documents", "/api/inventory", "/api/actions", "/api/settings", "/api/me"):
        assert client.get(path).status_code == 401
    assert client.get("/healthz").json()["status"] == "ok"


def test_full_journey_upload_review_approve_inventory_export(client: TestClient) -> None:
    _sign_up(client)

    pdf = _pdf("synthetic-001")
    results = _upload(client, ("Demowet NF SDS.pdf", pdf), ("inventory.xlsx", b"PK\x03\x04"))
    assert [r["accepted"] for r in results] == [True, False]
    assert results[1]["message"] == "Not a PDF. Only PDF files can be added."

    duplicate = _upload(client, ("again.pdf", pdf))  # same bytes, new name
    assert str(duplicate[0]["message"]).startswith("Already uploaded on")

    document = client.get("/api/documents").json()[0]
    assert document["status"] == "in_review" and document["open_reviews"] == 0

    review = client.get(f"/api/documents/{document['id']}/review").json()
    fields = {f["field_name"]: f for f in review["fields"]}
    assert fields["product_name"]["value"] == "Demowet NF"
    assert fields["product_name"]["status"] == "passed"
    assert fields["product_name"]["page"] == 1

    page = client.get(f"/api/documents/{document['id']}/pages/1").json()
    assert "Product name: Demowet NF" in page["text"]

    assert client.post(f"/api/documents/{document['id']}/approve").status_code == 200
    inventory = client.get("/api/inventory").json()
    assert inventory["total"] == 1
    product = inventory["products"][0]
    assert product["approved_by"] == "Rahima Akter"
    assert [h["h_code"] for h in product["h_codes"]] == ["H315", "H319"]
    assert {i["cas_number"] for i in product["ingredients"]} == {"67-63-0", "7732-18-5"}

    assert client.get("/api/inventory", params={"hazard": "H315"}).json()["products"]
    assert not client.get("/api/inventory", params={"q": "nothing like this"}).json()["products"]

    export = client.get("/api/export.xlsx")
    assert export.headers["content-disposition"].endswith('ChemReady_CIL_DemoWashingLtd_2026-10.xlsx"')
    workbook = load_workbook(io.BytesIO(export.content))
    assert workbook.sheetnames == ["CIL", "Flags"]
    rows = list(workbook["CIL"].values)
    assert rows[0][0] == "Product name" and len(rows) == 3  # header + one row per ingredient


def test_doubtful_fields_block_approval_until_resolved(client: TestClient) -> None:
    _sign_up(client)
    _upload(client, ("Demosoft SL SDS.pdf", _pdf("synthetic-004")))
    document = client.get("/api/documents").json()[0]
    assert document["open_reviews"] == 1  # no hazard statements: always reviewed

    blocked = client.post(f"/api/documents/{document['id']}/approve")
    assert blocked.status_code == 400 and "Resolve the 1 field" in blocked.json()["detail"]

    field = next(
        f
        for f in client.get(f"/api/documents/{document['id']}/review").json()["fields"]
        if f["status"] == "needs_review"
    )
    bad_edit = client.post(f"/api/fields/{field['id']}", json={"action": "edit", "value": "H3l9"})
    assert bad_edit.status_code == 400 and "H plus 3 digits" in bad_edit.json()["detail"]

    assert client.post(f"/api/fields/{field['id']}", json={"action": "missing"}).status_code == 200
    assert client.post(f"/api/documents/{document['id']}/approve").status_code == 200

    flags = client.get("/api/actions").json()
    assert {flag["kind"] for flag in flags} == {"missing_cas"}


def test_action_flags_follow_facility_settings(client: TestClient) -> None:
    _sign_up(client)
    _upload(client, ("Caustic.pdf", _pdf("synthetic-002")))  # revision date 2021-01-02
    document = client.get("/api/documents").json()[0]
    client.post(f"/api/documents/{document['id']}/approve")

    assert any(flag["kind"] == "old_sds" for flag in client.get("/api/actions").json())
    settings = client.get("/api/settings").json()
    client.put("/api/settings", json={**settings, "sds_max_age_years": 10})
    assert not any(flag["kind"] == "old_sds" for flag in client.get("/api/actions").json())


def test_unreadable_upload_becomes_a_failed_document_and_a_flag(client: TestClient) -> None:
    _sign_up(client)
    _upload(client, ("broken.pdf", b"%PDF-1.4 this is not really a pdf"))

    document = client.get("/api/documents").json()[0]
    assert document["status"] == "failed" and "Could not read this file" in document["error"]
    assert client.get("/api/actions").json()[0]["kind"] == "unreadable"
    assert client.post(f"/api/documents/{document['id']}/retry").status_code == 200
    assert client.delete(f"/api/documents/{document['id']}").status_code == 204
    assert client.get("/api/documents").json() == []


def test_one_facility_cannot_see_another_facilitys_documents(settings: Settings) -> None:
    app = create_app(settings, client_factory=RulesClient, inline_processing=True)
    first, second = TestClient(app), TestClient(app)
    _sign_up(first, "a@one.test", "Factory One")
    _sign_up(second, "b@two.test", "Factory Two")
    _upload(first, ("one.pdf", _pdf("synthetic-001")))
    document_id = first.get("/api/documents").json()[0]["id"]

    assert second.get("/api/documents").json() == []
    assert second.get(f"/api/documents/{document_id}/review").status_code == 404
    assert second.delete(f"/api/documents/{document_id}").status_code == 404
    field_id = first.get(f"/api/documents/{document_id}/review").json()["fields"][0]["id"]
    assert second.post(f"/api/fields/{field_id}", json={"action": "missing"}).status_code == 400


def test_private_file_is_refused_on_a_free_tier(settings: Settings) -> None:
    def free_tier() -> FakeClient:
        fake = FakeClient([])
        fake.allows_private_data = False
        return fake

    client = TestClient(create_app(settings, client_factory=free_tier, inline_processing=True))
    _sign_up(client)
    _upload(client, ("private.pdf", _pdf("synthetic-001")))

    document = client.get("/api/documents").json()[0]
    assert document["status"] == "failed" and "free tier" in document["error"]


def test_sign_in_wrong_password_and_lockout(client: TestClient) -> None:
    _sign_up(client)
    client.post("/api/auth/signout")

    for _ in range(auth.MAX_FAILED_SIGN_INS):
        wrong = client.post("/api/auth/signin", json={"email": "rahima@demo.test", "password": "nope-nope"})
        assert wrong.status_code == 401
    locked = client.post("/api/auth/signin", json={"email": "rahima@demo.test", "password": PASSWORD})
    assert "Too many attempts" in locked.json()["detail"]


def test_sign_up_validation_and_duplicate_email(client: TestClient) -> None:
    short = client.post(
        "/api/auth/signup", json={"name": "A", "email": "a@b.co", "password": "short", "facility_name": "F"}
    )
    assert short.status_code == 400 and "8 characters" in short.json()["detail"]
    _sign_up(client)
    again = client.post(
        "/api/auth/signup",
        json={"name": "R", "email": "RAHIMA@demo.test", "password": PASSWORD, "facility_name": "F"},
    )
    assert "already exists" in again.json()["detail"]


def test_password_reset_ends_old_sessions(client: TestClient) -> None:
    _sign_up(client)
    store: Store = client.app.state.chemready.store  # type: ignore[attr-defined]
    message = client.post("/api/auth/reset", json={"email": "nobody@nowhere.test"}).json()["message"]
    assert message.startswith("If an account exists")

    token = auth.start_password_reset(store, "rahima@demo.test")
    assert token is not None
    assert (
        client.post(
            "/api/auth/reset/confirm", json={"token": token, "new_password": "new password 123"}
        ).status_code
        == 200
    )
    assert client.get("/api/me").status_code == 401  # the old session was ended
    reused = client.post("/api/auth/reset/confirm", json={"token": token, "new_password": "another one 123"})
    assert reused.status_code == 400
    assert (
        client.post(
            "/api/auth/signin", json={"email": "rahima@demo.test", "password": "new password 123"}
        ).status_code
        == 200
    )


def test_sign_out_everywhere(client: TestClient) -> None:
    _sign_up(client)
    other_device = TestClient(client.app)
    assert (
        other_device.post(
            "/api/auth/signin", json={"email": "rahima@demo.test", "password": PASSWORD}
        ).status_code
        == 200
    )

    client.post("/api/auth/signout-all")

    assert other_device.get("/api/me").status_code == 401


def test_background_worker_processes_queued_documents(settings: Settings, tmp_path: Path) -> None:
    app = create_app(settings, client_factory=RulesClient)
    with TestClient(app) as client:  # runs the start-up hook, which starts the worker thread
        _sign_up(client)
        _upload(client, ("one.pdf", _pdf("synthetic-001")))
        for _ in range(100):
            if client.get("/api/documents").json()[0]["status"] == "in_review":
                break
            time.sleep(0.05)
        assert client.get("/api/documents").json()[0]["status"] == "in_review"


def test_worker_picks_up_documents_left_from_a_restart(settings: Settings) -> None:
    store = Store(settings.database_path)
    facility = store.create_facility("F", 3)
    user = store.create_user(facility, "U", "u@f.test", auth.hash_password(PASSWORD))
    path = settings.upload_dir / "left.pdf"
    path.parent.mkdir(parents=True)
    path.write_bytes(_pdf("synthetic-001"))
    document_id = store.add_document(facility, "left.pdf", "hash", str(path), user, True)

    worker = Worker(store, RulesClient)
    worker.start()
    worker.stop()

    assert store.document(facility, document_id)["status"] == "in_review"  # type: ignore[index]


def test_seed_demo_creates_account_with_reviewed_and_waiting_files(settings: Settings) -> None:
    from chemready.demo.seed import DEMO_EMAIL, DEMO_PASSWORD, seed_demo

    store = Store(settings.database_path)
    assert "created" in seed_demo(settings, store)
    assert "already exists" in seed_demo(settings, store)

    client = TestClient(create_app(settings, store=store, client_factory=RulesClient, inline_processing=True))
    assert (
        client.post("/api/auth/signin", json={"email": DEMO_EMAIL, "password": DEMO_PASSWORD}).status_code
        == 200
    )
    statuses = sorted(d["status"] for d in client.get("/api/documents").json())
    assert statuses == ["approved", "approved", "in_review", "in_review"]


def test_cli_commands(
    monkeypatch: pytest.MonkeyPatch, settings: Settings, capsys: pytest.CaptureFixture[str]
) -> None:
    import chemready.__main__ as cli

    monkeypatch.setattr(cli, "get_settings", lambda: settings)
    cli.main(["seed-demo"])
    assert "Demo account created" in capsys.readouterr().out

    started: dict[str, object] = {}
    monkeypatch.setattr("uvicorn.run", lambda *a, **k: started.update(k))
    cli.main(["serve", "--port", "9999"])
    assert started["port"] == 9999
